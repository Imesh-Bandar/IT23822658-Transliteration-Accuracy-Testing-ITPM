"""
Generates  IT23822658 - Assignment 1 - Test cases.xlsx
Run:  python generate_excel.py
Requires:  pip install openpyxl
"""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from test_cases_data import TEST_CASES

HEADER_FILL   = PatternFill("solid", fgColor="1F3864")
FAIL_FILL     = PatternFill("solid", fgColor="FFCCCC")
ALT_ROW_FILL  = PatternFill("solid", fgColor="F2F2F2")
WHITE_FILL    = PatternFill("solid", fgColor="FFFFFF")

HEADER_FONT   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT     = Font(name="Calibri", size=10)
FAIL_FONT     = Font(name="Calibri", bold=True, color="C00000", size=10)
WRAP          = Alignment(wrap_text=True, vertical="top")
CENTER_WRAP   = Alignment(wrap_text=True, vertical="top", horizontal="center")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("Test Case ID",                 14),
    ("Input\nlength type",           10),
    ("Input",                        38),
    ("Expected output",              38),
    ("Actual output",                38),
    ("Status",                       10),
    ("Singlish input types covered", 32),
    ("Evidence / Rationale",         52),
]

def make_cell(ws, row, col, value, font=None, fill=None, alignment=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:      cell.font      = font
    if fill:      cell.fill      = fill
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border
    return cell


def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Negative Test Cases"
    ws.sheet_view.showGridLines = False

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        make_cell(
            ws, 1, col_idx, header,
            font=HEADER_FONT,
            fill=HEADER_FILL,
            alignment=CENTER_WRAP,
            border=BORDER,
        )
    ws.row_dimensions[1].height = 30

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, tc in enumerate(TEST_CASES, start=2):
        fill = FAIL_FILL if row_idx % 2 == 0 else ALT_ROW_FILL

        values = [
            tc["id"],
            tc["length_type"],
            tc["input"],
            tc["expected"],
            tc["actual"],
            tc["status"],
            tc["types"],
            tc["rationale"],
        ]

        for col_idx, value in enumerate(values, start=1):
            is_status = (col_idx == 6)
            is_id     = (col_idx == 1)
            f = FAIL_FONT if is_status else BODY_FONT
            a = CENTER_WRAP if (is_status or is_id) else WRAP
            make_cell(ws, row_idx, col_idx, value, font=f, fill=fill,
                      alignment=a, border=BORDER)

        ws.row_dimensions[row_idx].height = 80

    # ── Freeze header ─────────────────────────────────────────────────────────
    ws.freeze_panes = "A2"

    filename = "IT23822658 - Assignment 1 - Test cases.xlsx"
    wb.save(filename)
    print(f"Saved → {filename}  ({len(TEST_CASES)} test cases)")


if __name__ == "__main__":
    build_workbook()
